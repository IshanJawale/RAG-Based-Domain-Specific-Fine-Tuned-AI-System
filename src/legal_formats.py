"""
legal_formats.py - Registry of Indian legal document formats and a generic .docx renderer.

Each format declares:
  label          : human-readable name
  basis          : statutory basis shown in the UI and used as a retrieval hint
  patterns       : regexes that sub-classify a drafting request into this format (checked in FORMAT_ORDER)
  fields         : field name -> description (given to the LLM as the JSON schema to fill)
  layout         : ordered blocks rendered into the document (see _render_block)
  notes          : drafter's checklist printed at the end ("remove before filing")

Missing values are rendered as "[●]", the usual blank in Indian drafting, never invented.
Statutory references follow BNSS/BNS/BSA 2023 with the old CrPC/IPC section in brackets where useful.
The checklists are drafting aids, not legal advice: verify against the current bare act and court rules.
"""

import re
import string
from datetime import datetime
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, Mm, RGBColor

DOCS_DIR = Path("generated_documents")
BLANK = "[●]"

# ----------------------------------------------------------------------------
# Shared field groups
# ----------------------------------------------------------------------------
_NOTICE_PARTIES = {
    "sender_name": "Client on whose behalf the notice is sent",
    "sender_address": "Client's address",
    "recipient_name": "Person or company the notice is addressed to",
    "recipient_address": "Recipient's address",
    "advocate_name": "Advocate issuing the notice (name, enrolment no.)",
    "advocate_address": "Advocate's chamber address",
}
_SUIT_PARTIES = {
    "court_name": "Court, e.g. 'Court of the Civil Judge (Senior Division), Pune'",
    "case_number": "Case number line, e.g. 'Civil Suit No. ____ of 2026'",
    "petitioner_name": "Plaintiff / applicant / petitioner name with parentage, age, occupation",
    "petitioner_address": "Plaintiff / applicant / petitioner address",
    "respondent_name": "Defendant / respondent name",
    "respondent_address": "Defendant / respondent address",
    "place": "Place of signing",
}

_VERIFY_PLEADING = (
    "VERIFICATION\n\nI, {verifier}, the {party_role} above named, do hereby verify that the contents of "
    "paragraphs 1 to {last_para} of the above are true to my knowledge, and those of the remaining paragraphs are "
    "based on legal advice which I believe to be true. No part of it is false and nothing material has been concealed.\n\n"
    "Verified at {place} on this {date}."
)

FORMATS: dict[str, dict] = {
    # ------------------------------------------------------------------ notices
    "cheque_bounce_notice": {
        "label": "Cheque dishonour notice (Section 138 NI Act)",
        "category": "notice",
        "basis": "Sections 138 and 142, Negotiable Instruments Act, 1881",
        "patterns": [r"cheque\s*(?:bounce|dishono?u?r|return)", r"\b138\b.*\b(?:ni|negotiable)", r"dishono?u?r(?:ed)?\s+(?:of\s+)?cheque"],
        "fields": {
            **_NOTICE_PARTIES,
            "liability": "Legally enforceable debt or liability for which the cheque was issued",
            "cheque_number": "Cheque number",
            "cheque_date": "Date on the cheque",
            "cheque_amount": "Cheque amount in Rs.",
            "bank_name": "Drawee bank and branch",
            "dishonour_reason": "Reason on the return memo, e.g. 'Funds Insufficient'",
            "return_memo_date": "Date of the bank's return memo",
            "facts": "Any additional facts (list)",
        },
        "layout": [
            {"type": "banner", "text": "BY REGISTERED POST A.D. / SPEED POST\nLEGAL NOTICE UNDER SECTION 138 OF THE NEGOTIABLE INSTRUMENTS ACT, 1881"},
            {"type": "line", "text": "Date: {date}", "align": "right"},
            {"type": "line", "text": "To,\n{recipient_name}\n{recipient_address}"},
            {"type": "line", "text": "Sub: Legal notice for dishonour of cheque no. {cheque_number} dated {cheque_date} for Rs. {cheque_amount} drawn on {bank_name}", "bold": True},
            {"type": "line", "text": "Sir/Madam,\n\nUnder instructions from and on behalf of my client, {sender_name}, residing at {sender_address}, I serve upon you the following notice:"},
            {"type": "numbered", "items": [
                "That you were liable to pay my client towards {liability}.",
                "That in discharge of the said legally enforceable debt/liability, you issued cheque no. {cheque_number} dated {cheque_date} for Rs. {cheque_amount} drawn on {bank_name} in favour of my client.",
                "That on presentation, the said cheque was returned unpaid with the remark \"{dishonour_reason}\" vide return memo dated {return_memo_date}.",
            ]},
            {"type": "numbered", "field": "facts"},
            {"type": "numbered", "items": [
                "That the dishonour of the said cheque constitutes an offence punishable under Section 138 of the Negotiable Instruments Act, 1881.",
                "I therefore call upon you to pay the said sum of Rs. {cheque_amount} to my client within 15 (fifteen) days of receipt of this notice, failing which my client shall be constrained to file a complaint under Sections 138 and 142 of the Negotiable Instruments Act, 1881, besides civil proceedings for recovery, entirely at your risk as to costs and consequences.",
            ]},
            {"type": "signature", "text": "{advocate_name}\nAdvocate for {sender_name}\n{advocate_address}"},
        ],
        "notes": [
            "Send within 30 days of receiving the bank's return memo (Section 138(b)).",
            "The drawer has 15 days from receipt to pay; the complaint lies only after that period expires (Yogendra Pratap Singh v. Savitri Pandey, 2014).",
            "File the complaint within one month of the cause of action, i.e. after the 15 days lapse (Section 142(1)(b)).",
            "Keep postal receipts and tracking/A.D. card as proof of service.",
        ],
    },
    "section_80_cpc_notice": {
        "label": "Notice to Government (Section 80 CPC)",
        "category": "notice",
        "basis": "Section 80, Code of Civil Procedure, 1908",
        "patterns": [r"section\s*80\b", r"\b80\s*cpc", r"notice\s+(?:to|against)\s+(?:the\s+)?(?:state|central|union)?\s*govern?ment", r"suit\s+against\s+(?:the\s+)?govern?ment"],
        "fields": {
            **_NOTICE_PARTIES,
            "recipient_name": "Officer to be served, e.g. 'The Secretary to the Government of Maharashtra, Revenue Department' or 'The Collector, Pune'",
            "subject": "Subject line",
            "facts": "Facts giving rise to the claim (list)",
            "cause_of_action": "When and how the cause of action arose",
            "claims": "Relief the client will claim in the suit (list)",
        },
        "layout": [
            {"type": "banner", "text": "BY REGISTERED POST A.D.\nNOTICE UNDER SECTION 80 OF THE CODE OF CIVIL PROCEDURE, 1908"},
            {"type": "line", "text": "Date: {date}", "align": "right"},
            {"type": "line", "text": "To,\n{recipient_name}\n{recipient_address}"},
            {"type": "line", "text": "Sub: {subject}", "bold": True},
            {"type": "line", "text": "Sir,\n\nUnder instructions from my client, {sender_name}, residing at {sender_address}, I give you notice under Section 80 of the Code of Civil Procedure, 1908, as follows:"},
            {"type": "heading", "text": "Facts"},
            {"type": "numbered", "field": "facts"},
            {"type": "heading", "text": "Cause of action"},
            {"type": "numbered", "field": "cause_of_action"},
            {"type": "heading", "text": "Relief claimed"},
            {"type": "numbered", "field": "claims"},
            {"type": "line", "text": "Take notice that if the above relief is not granted within two months of delivery of this notice, my client shall institute a suit against the Government for the said relief, with costs, without any further notice."},
            {"type": "signature", "text": "{advocate_name}\nAdvocate for {sender_name}\n{advocate_address}"},
        ],
        "notes": [
            "A suit against the Government or a public officer cannot be filed until two months after this notice is delivered (Section 80(1)).",
            "For urgent relief, a suit may be filed without notice with the court's leave (Section 80(2)).",
            "The notice must state the name, description and residence of the claimant, the cause of action and the relief claimed.",
        ],
    },
    "legal_notice": {
        "label": "Legal notice (demand notice)",
        "category": "notice",
        "basis": "General demand notice; cite the provisions breached",
        "patterns": [r"\bnotice\b"],
        "fields": {
            **_NOTICE_PARTIES,
            "subject": "Subject line",
            "facts": "Facts in chronological order (list)",
            "legal_provisions": "Provisions breached, e.g. 'Section 73, Indian Contract Act, 1872' (list)",
            "claims": "Demands: amounts, actions, costs (list)",
            "demand_period_days": "Days given to comply (default 15)",
        },
        "layout": [
            {"type": "banner", "text": "BY REGISTERED POST A.D. / SPEED POST\nLEGAL NOTICE"},
            {"type": "line", "text": "Date: {date}", "align": "right"},
            {"type": "line", "text": "To,\n{recipient_name}\n{recipient_address}"},
            {"type": "line", "text": "Sub: {subject}", "bold": True},
            {"type": "line", "text": "Sir/Madam,\n\nUnder instructions from and on behalf of my client, {sender_name}, residing at {sender_address}, I serve upon you the following legal notice:"},
            {"type": "numbered", "field": "facts"},
            {"type": "heading", "text": "Legal provisions"},
            {"type": "bullets", "field": "legal_provisions"},
            {"type": "heading", "text": "Demand"},
            {"type": "bullets", "field": "claims"},
            {"type": "line", "text": "You are hereby called upon to comply with the above within {demand_period_days} days of receipt of this notice, failing which my client shall initiate appropriate civil and/or criminal proceedings against you, entirely at your risk as to costs and consequences."},
            {"type": "line", "text": "A copy of this notice has been retained in my office for record and further action."},
            {"type": "signature", "text": "{advocate_name}\nAdvocate for {sender_name}\n{advocate_address}"},
        ],
        "notes": [
            "Send by registered post A.D. or speed post and keep the receipts.",
            "If the recipient is the Government or a public officer, use the Section 80 CPC notice instead.",
            "Check the limitation period for the eventual suit before giving a long compliance period.",
        ],
    },
    # ---------------------------------------------------------------- pleadings
    "plaint": {
        "label": "Plaint (civil suit, Order VII CPC)",
        "category": "pleading",
        "basis": "Order VII Rule 1 and Order VI Rule 15, Code of Civil Procedure, 1908",
        "patterns": [r"\bplaint\b", r"civil\s+suit", r"suit\s+for\s+(?:recovery|declaration|injunction|possession|partition|specific\s+performance|damages)", r"recovery\s+suit"],
        "fields": {
            **_SUIT_PARTIES,
            "suit_title": "Nature of the suit, e.g. 'Suit for recovery of Rs. 15,00,000 with interest'",
            "facts": "Facts constituting the cause of action, in numbered paragraphs (list)",
            "cause_of_action": "When the cause of action arose and continues",
            "jurisdiction": "Territorial and pecuniary jurisdiction of this court",
            "limitation": "Why the suit is within limitation",
            "valuation": "Valuation for jurisdiction and court fee paid",
            "prayer": "Reliefs prayed for, each separately (list)",
        },
        "roles": ("Plaintiff", "Defendant"),
        "layout": [
            {"type": "cause_title"},
            {"type": "title", "text": "{suit_title}"},
            {"type": "line", "text": "The Plaintiff above named most respectfully submits as under:"},
            {"type": "numbered", "field": "facts"},
            {"type": "heading", "text": "Cause of action"},
            {"type": "numbered", "field": "cause_of_action"},
            {"type": "heading", "text": "Jurisdiction"},
            {"type": "numbered", "field": "jurisdiction"},
            {"type": "heading", "text": "Limitation"},
            {"type": "numbered", "field": "limitation"},
            {"type": "heading", "text": "Valuation and court fee"},
            {"type": "numbered", "field": "valuation"},
            {"type": "prayer", "intro": "It is therefore most respectfully prayed that this Hon'ble Court may be pleased to pass a decree in favour of the Plaintiff and against the Defendant:"},
            {"type": "signature", "text": "PLAINTIFF\n\nThrough Counsel\n\nPlace: {place}\nDate: {date}"},
            {"type": "verification", "text": _VERIFY_PLEADING},
        ],
        "notes": [
            "Contents required by Order VII Rule 1: court, parties, facts of the cause of action and when it arose, jurisdiction, relief, valuation and court fee.",
            "Verify under Order VI Rule 15 and file the supporting affidavit; commercial suits need a statement of truth.",
            "List and file the documents relied on (Order VII Rule 14). Pay court fee under the State Court Fees Act.",
            "Commercial disputes: pre-institution mediation under Section 12A, Commercial Courts Act, 2015, unless urgent relief is sought.",
        ],
    },
    "written_statement": {
        "label": "Written statement (Order VIII CPC)",
        "category": "pleading",
        "basis": "Order VIII, Code of Civil Procedure, 1908",
        "patterns": [r"written\s+statement", r"reply\s+to\s+(?:the\s+)?plaint", r"defen[cs]e\s+(?:to|in)\s+(?:the\s+)?suit"],
        "fields": {
            **_SUIT_PARTIES,
            "preliminary_objections": "Preliminary objections: maintainability, limitation, jurisdiction, cause of action (list)",
            "parawise_reply": "Para-wise reply to the plaint, e.g. 'Para 1: denied...' (list)",
            "additional_facts": "Defendant's own version of facts (list)",
            "prayer": "Prayer, e.g. dismissal with costs (list)",
        },
        "roles": ("Plaintiff", "Defendant"),
        "verify_as": "Defendant",
        "layout": [
            {"type": "cause_title"},
            {"type": "title", "text": "WRITTEN STATEMENT ON BEHALF OF THE DEFENDANT"},
            {"type": "line", "text": "The Defendant above named most respectfully submits as under:"},
            {"type": "heading", "text": "Preliminary objections"},
            {"type": "numbered", "field": "preliminary_objections"},
            {"type": "heading", "text": "Reply on merits (para-wise)"},
            {"type": "numbered", "field": "parawise_reply"},
            {"type": "heading", "text": "Additional facts"},
            {"type": "numbered", "field": "additional_facts"},
            {"type": "prayer", "intro": "It is therefore most respectfully prayed that this Hon'ble Court may be pleased to:"},
            {"type": "signature", "text": "DEFENDANT\n\nThrough Counsel\n\nPlace: {place}\nDate: {date}"},
            {"type": "verification", "text": _VERIFY_PLEADING},
        ],
        "notes": [
            "File within 30 days of service of summons, extendable to 90 days (Order VIII Rule 1); in commercial suits 120 days is the outer limit.",
            "Deny each allegation specifically; an evasive denial is treated as an admission (Order VIII Rules 3 to 5).",
            "Raise set-off or counter-claim, if any, in the written statement itself (Order VIII Rules 6 and 6A).",
        ],
    },
    # ---------------------------------------------------------------- criminal
    "anticipatory_bail_application": {
        "label": "Anticipatory bail application (Section 482 BNSS)",
        "category": "criminal",
        "basis": "Section 482, Bharatiya Nagarik Suraksha Sanhita, 2023 (formerly Section 438 CrPC)",
        "patterns": [r"anticipatory\s+bail", r"pre[- ]?arrest\s+bail", r"\b482\b.*bnss", r"\b438\b.*cr\.?p\.?c"],
        "fields": {
            **_SUIT_PARTIES,
            "court_name": "Court of Session or High Court",
            "case_number": "e.g. 'Criminal Misc. Bail Application No. ____ of 2026'",
            "respondent_name": "Usually 'State of <State>' through the police station",
            "fir_number": "FIR number and year (or 'apprehended FIR')",
            "police_station": "Police station",
            "offences": "Sections alleged, e.g. 'Sections 318(4), 316(2) BNS'",
            "facts": "Brief facts and the reason arrest is apprehended (list)",
            "grounds": "Grounds for anticipatory bail (list)",
            "previous_applications": "Whether any earlier bail application was filed and its result",
            "prayer": "Prayer (list)",
        },
        "roles": ("Applicant", "Respondent"),
        "verify_as": "Applicant",
        "layout": [
            {"type": "cause_title"},
            {"type": "title", "text": "APPLICATION UNDER SECTION 482 OF THE BHARATIYA NAGARIK SURAKSHA SANHITA, 2023 FOR GRANT OF ANTICIPATORY BAIL"},
            {"type": "line", "text": "FIR No.: {fir_number}    Police Station: {police_station}    Offences: {offences}", "bold": True},
            {"type": "line", "text": "The Applicant most respectfully submits as under:"},
            {"type": "numbered", "field": "facts"},
            {"type": "heading", "text": "Grounds"},
            {"type": "lettered", "field": "grounds"},
            {"type": "numbered", "items": ["That the details of any previous bail application filed by the Applicant in this case are: {previous_applications}.", "That the Applicant undertakes to abide by all conditions under Section 482(2) of the BNSS, including joining the investigation as and when required, not leaving India without the Court's permission and not influencing witnesses."]},
            {"type": "prayer", "intro": "It is therefore most respectfully prayed that this Hon'ble Court may be pleased to:"},
            {"type": "signature", "text": "APPLICANT\n\nThrough Counsel\n\nPlace: {place}\nDate: {date}"},
            {"type": "verification", "text": _VERIFY_PLEADING},
        ],
        "notes": [
            "Lies before the Court of Session or the High Court (concurrent jurisdiction); practice usually requires approaching Sessions first.",
            "Disclose all earlier bail applications and their outcome.",
            "Annex the FIR or complaint (if registered), any notice received, and an affidavit in support.",
        ],
    },
    "bail_application": {
        "label": "Regular bail application (Sections 480/483 BNSS)",
        "category": "criminal",
        "basis": "Sections 480 and 483, Bharatiya Nagarik Suraksha Sanhita, 2023 (formerly Sections 437/439 CrPC)",
        "patterns": [r"\bbail\b"],
        "fields": {
            **_SUIT_PARTIES,
            "court_name": "Magistrate (Section 480) or Sessions Court / High Court (Section 483)",
            "case_number": "e.g. 'Bail Application No. ____ of 2026'",
            "respondent_name": "Usually 'State of <State>'",
            "fir_number": "FIR number and year",
            "police_station": "Police station",
            "offences": "Sections alleged",
            "date_of_arrest": "Date of arrest / custody",
            "facts": "Brief facts and prosecution case (list)",
            "grounds": "Grounds for bail: parity, custody period, no flight risk, investigation complete, etc. (list)",
            "previous_applications": "Earlier bail applications and their result",
            "prayer": "Prayer (list)",
        },
        "roles": ("Applicant/Accused", "Respondent"),
        "verify_as": "Applicant",
        "layout": [
            {"type": "cause_title"},
            {"type": "title", "text": "APPLICATION UNDER SECTION 480 / 483 OF THE BHARATIYA NAGARIK SURAKSHA SANHITA, 2023 FOR GRANT OF REGULAR BAIL"},
            {"type": "line", "text": "FIR No.: {fir_number}    Police Station: {police_station}    Offences: {offences}    In custody since: {date_of_arrest}", "bold": True},
            {"type": "line", "text": "The Applicant most respectfully submits as under:"},
            {"type": "numbered", "field": "facts"},
            {"type": "heading", "text": "Grounds"},
            {"type": "lettered", "field": "grounds"},
            {"type": "numbered", "items": ["That the details of any previous bail application filed by the Applicant in this case are: {previous_applications}.", "That the Applicant undertakes to abide by any condition imposed by this Hon'ble Court and to cooperate with the trial."]},
            {"type": "prayer", "intro": "It is therefore most respectfully prayed that this Hon'ble Court may be pleased to:"},
            {"type": "signature", "text": "APPLICANT\n\nThrough Counsel\n\nPlace: {place}\nDate: {date}"},
            {"type": "verification", "text": _VERIFY_PLEADING},
        ],
        "notes": [
            "Section 480 BNSS: Magistrate. Section 483 BNSS: Sessions Court or High Court, including after refusal by the Magistrate.",
            "Check default bail under Section 187(3) BNSS (formerly Section 167(2) CrPC) if the charge-sheet was not filed within 60/90 days, and undertrial release under Section 479 BNSS.",
            "Annex the FIR, arrest memo, remand orders and the order on any earlier bail application.",
        ],
    },
    "police_complaint": {
        "label": "Police complaint / FIR request (Section 173 BNSS)",
        "category": "criminal",
        "basis": "Section 173, Bharatiya Nagarik Suraksha Sanhita, 2023 (formerly Section 154 CrPC)",
        "patterns": [r"\bfir\b", r"police\s+complaint", r"complaint\s+to\s+(?:the\s+)?police", r"\b173\b.*bnss", r"\b154\b.*cr\.?p\.?c"],
        "fields": {
            "police_station": "Police station and the SHO's designation",
            "petitioner_name": "Complainant's name and parentage",
            "petitioner_address": "Complainant's address and phone",
            "respondent_name": "Accused person(s), or 'unknown persons'",
            "incident_date": "Date and time of the incident",
            "incident_place": "Place of the incident",
            "facts": "What happened, in order (list)",
            "offences": "Offences disclosed, e.g. 'Sections 303(2), 351 BNS'",
            "place": "Place",
        },
        "layout": [
            {"type": "line", "text": "To,\nThe Station House Officer,\n{police_station}"},
            {"type": "line", "text": "Date: {date}", "align": "right"},
            {"type": "line", "text": "Sub: Complaint for registration of FIR under Section 173 of the Bharatiya Nagarik Suraksha Sanhita, 2023", "bold": True},
            {"type": "line", "text": "Sir/Madam,\n\nI, {petitioner_name}, residing at {petitioner_address}, submit this complaint against {respondent_name}. On {incident_date} at {incident_place}, the following took place:"},
            {"type": "numbered", "field": "facts"},
            {"type": "line", "text": "The above acts disclose cognizable offences including {offences}. I request you to register an FIR, investigate the matter and take action in accordance with law, and to provide me a free copy of the FIR as required by Section 173(2) BNSS."},
            {"type": "signature", "text": "Yours faithfully,\n\n{petitioner_name}\nPlace: {place}"},
        ],
        "notes": [
            "Obtain a receipt / diary number for the complaint.",
            "If the SHO refuses to register, send the substance to the Superintendent of Police (Section 173(4) BNSS), then approach the Magistrate under Section 175(3) BNSS.",
            "For offences punishable with 3 to 7 years, the officer may hold a preliminary enquiry within 14 days (Section 173(3) BNSS).",
        ],
    },
    # -------------------------------------------------------------- petitions
    "writ_petition": {
        "label": "Writ petition (Article 226)",
        "category": "petition",
        "basis": "Article 226, Constitution of India",
        "patterns": [r"\bwrit\b", r"article\s*226", r"\bmandamus\b", r"\bcertiorari\b", r"habeas\s+corpus"],
        "fields": {
            **_SUIT_PARTIES,
            "court_name": "High Court, e.g. 'High Court of Judicature at Bombay'",
            "case_number": "e.g. 'Writ Petition (Civil) No. ____ of 2026'",
            "respondent_name": "Respondents, e.g. 'Union of India through the Secretary, ...' (list allowed)",
            "synopsis": "Short synopsis of the case",
            "list_of_dates": "Chronology, each item 'DD.MM.YYYY: event' (list)",
            "facts": "Facts (list)",
            "questions_of_law": "Questions of law (list)",
            "grounds": "Grounds of challenge (list)",
            "alternate_remedy": "Why no efficacious alternative remedy is available",
            "interim_relief": "Interim relief sought (list)",
            "prayer": "Final reliefs, e.g. writ of certiorari quashing order dated ... (list)",
        },
        "roles": ("Petitioner", "Respondent"),
        "verify_as": "Petitioner",
        "layout": [
            {"type": "cause_title", "extra": "(EXTRAORDINARY CIVIL WRIT JURISDICTION)"},
            {"type": "title", "text": "WRIT PETITION UNDER ARTICLE 226 OF THE CONSTITUTION OF INDIA"},
            {"type": "heading", "text": "Synopsis"},
            {"type": "text", "field": "synopsis"},
            {"type": "heading", "text": "List of dates and events"},
            {"type": "bullets", "field": "list_of_dates"},
            {"type": "line", "text": "To,\nThe Hon'ble Chief Justice and His/Her Companion Justices of the {court_name}.\n\nThe humble petition of the Petitioner above named most respectfully showeth:"},
            {"type": "numbered", "field": "facts"},
            {"type": "heading", "text": "Questions of law"},
            {"type": "numbered", "field": "questions_of_law"},
            {"type": "heading", "text": "Grounds"},
            {"type": "lettered", "field": "grounds"},
            {"type": "numbered", "items": ["{alternate_remedy}", "That the Petitioner has not filed any other petition before this Hon'ble Court or the Supreme Court on the same cause of action."]},
            {"type": "heading", "text": "Interim relief"},
            {"type": "bullets", "field": "interim_relief"},
            {"type": "prayer", "intro": "It is therefore most respectfully prayed that this Hon'ble Court may be pleased to:"},
            {"type": "signature", "text": "PETITIONER\n\nThrough Counsel\n\nPlace: {place}\nDate: {date}"},
            {"type": "verification", "text": _VERIFY_PLEADING},
        ],
        "notes": [
            "Annex the impugned order/action as Annexure P-1 and other documents serially; file a supporting affidavit.",
            "Explain why the alternative statutory remedy (appeal/revision/tribunal) is not efficacious, or the petition may be dismissed on that ground.",
            "Follow the specific High Court's Writ Rules for index, court fee, urgency application and number of copies.",
        ],
    },
    "consumer_complaint": {
        "label": "Consumer complaint (Section 35 CPA 2019)",
        "category": "petition",
        "basis": "Sections 35 and 39, Consumer Protection Act, 2019",
        "patterns": [r"consumer\s+(?:complaint|forum|court|commission|case)", r"deficiency\s+in\s+service", r"consumer\s+protection\s+act", r"defective\s+(?:product|goods)"],
        "fields": {
            **_SUIT_PARTIES,
            "court_name": "e.g. 'District Consumer Disputes Redressal Commission, Pune'",
            "case_number": "e.g. 'Consumer Complaint No. ____ of 2026'",
            "respondent_name": "Opposite party (seller / service provider)",
            "consideration_paid": "Price or consideration paid, Rs.",
            "facts": "Facts: purchase/service, defect or deficiency, complaints made (list)",
            "deficiency": "The defect / deficiency in service / unfair trade practice alleged",
            "jurisdiction": "Pecuniary and territorial jurisdiction of this Commission",
            "limitation": "Date the cause of action arose; filed within 2 years",
            "prayer": "Reliefs under Section 39: refund, replacement, compensation, costs (list)",
        },
        "roles": ("Complainant", "Opposite Party"),
        "verify_as": "Complainant",
        "layout": [
            {"type": "cause_title"},
            {"type": "title", "text": "COMPLAINT UNDER SECTION 35 OF THE CONSUMER PROTECTION ACT, 2019"},
            {"type": "line", "text": "The Complainant most respectfully submits as under:"},
            {"type": "numbered", "field": "facts"},
            {"type": "numbered", "items": ["That the Opposite Party is guilty of {deficiency}.", "That the Complainant paid a consideration of Rs. {consideration_paid}."]},
            {"type": "heading", "text": "Jurisdiction"},
            {"type": "numbered", "field": "jurisdiction"},
            {"type": "heading", "text": "Limitation"},
            {"type": "numbered", "field": "limitation"},
            {"type": "prayer", "intro": "It is therefore most respectfully prayed that this Hon'ble Commission may be pleased to direct the Opposite Party to:"},
            {"type": "signature", "text": "COMPLAINANT\n\nPlace: {place}\nDate: {date}"},
            {"type": "verification", "text": _VERIFY_PLEADING},
        ],
        "notes": [
            "Pecuniary jurisdiction is by consideration paid: District up to Rs. 50 lakh, State Rs. 50 lakh to 2 crore, National above Rs. 2 crore (2021 Jurisdiction Rules); verify current limits.",
            "File within 2 years of the cause of action (Section 69), or with an application for condonation of delay.",
            "Can be filed online on e-Daakhil; attach invoice, correspondence, and proof of payment of the prescribed fee.",
        ],
    },
    "rti_application": {
        "label": "RTI application (Section 6(1) RTI Act)",
        "category": "application",
        "basis": "Section 6(1), Right to Information Act, 2005",
        "patterns": [r"\brti\b", r"right\s+to\s+information"],
        "fields": {
            "pio_designation": "Public Information Officer's designation",
            "public_authority": "Public authority / department",
            "authority_address": "Address of the public authority",
            "petitioner_name": "Applicant's name",
            "petitioner_address": "Applicant's address for reply",
            "information_sought": "Specific information requested, one point per item (list)",
            "period": "Period the information relates to",
            "fee_details": "Fee payment details, e.g. 'Rs. 10 by IPO No. ____' or 'BPL - exempt'",
            "place": "Place",
        },
        "layout": [
            {"type": "line", "text": "To,\nThe Public Information Officer,\n{pio_designation}\n{public_authority}\n{authority_address}"},
            {"type": "line", "text": "Date: {date}", "align": "right"},
            {"type": "line", "text": "Sub: Application for information under Section 6(1) of the Right to Information Act, 2005", "bold": True},
            {"type": "line", "text": "Sir/Madam,\n\nI, {petitioner_name}, a citizen of India, request the following information relating to the period {period}:"},
            {"type": "numbered", "field": "information_sought"},
            {"type": "line", "text": "Application fee: {fee_details}.\n\nIf the information or part of it is held by another public authority, kindly transfer this application under Section 6(3). I would prefer copies of documents and records wherever applicable."},
            {"type": "signature", "text": "Yours faithfully,\n\n{petitioner_name}\n{petitioner_address}\nPlace: {place}"},
        ],
        "notes": [
            "Central fee is Rs. 10 (IPO, DD, banker's cheque, cash or online); State public authorities follow State RTI rules.",
            "Below-poverty-line applicants are exempt on proof (Section 7(5)).",
            "Reply due within 30 days, or 48 hours where life or liberty is involved (Section 7(1)); first appeal within 30 days (Section 19(1)).",
            "Ask for specific records rather than opinions or reasons; no reason for seeking information is required (Section 6(2)).",
        ],
    },
    "affidavit": {
        "label": "Affidavit",
        "category": "affidavit",
        "basis": "General affidavit; Order XIX CPC where filed in court",
        "patterns": [r"\baffidavit\b"],
        "fields": {
            "court_name": "Court, if the affidavit is to be filed in a case (optional)",
            "case_number": "Case number, if any (optional)",
            "petitioner_name": "Deponent's full name and parentage",
            "deponent_details": "Age, occupation and address of the deponent",
            "purpose": "Purpose of the affidavit",
            "statements": "Statements on oath, one fact per item (list)",
            "place": "Place",
        },
        "layout": [
            {"type": "line", "text": "{court_name}\n{case_number}", "align": "center", "bold": True, "optional": True},
            {"type": "title", "text": "AFFIDAVIT"},
            {"type": "line", "text": "I, {petitioner_name}, {deponent_details}, do hereby solemnly affirm and state on oath as under, for the purpose of {purpose}:"},
            {"type": "numbered", "field": "statements"},
            {"type": "signature", "text": "DEPONENT"},
            {"type": "line", "text": "VERIFICATION\n\nVerified at {place} on this {date} that the contents of the above affidavit are true and correct to my knowledge and belief, no part of it is false and nothing material has been concealed therefrom."},
            {"type": "signature", "text": "DEPONENT"},
        ],
        "notes": [
            "Execute on non-judicial stamp paper of the value prescribed by the State Stamp Act, where required.",
            "Sign before a Notary Public or Oath Commissioner, who attests and seals it.",
            "Keep statements to facts within the deponent's knowledge; distinguish knowledge from belief (Order XIX Rule 3 CPC).",
        ],
    },
}

# Advisory formats rendered by the dedicated functions in document_generator.py
ADVISORY_FORMATS = {
    "legal_opinion": {"label": "Legal opinion / memorandum", "category": "advisory", "basis": "Advisory opinion", "patterns": [r"\bopinion\b", r"\bmemo(?:randum)?\b"]},
    "case_brief": {"label": "Case brief", "category": "advisory", "basis": "Judgment analysis", "patterns": [r"case\s+brief", r"\bbrief\b"]},
    "legal_report": {"label": "Legal assessment report", "category": "advisory", "basis": "Research / advisory report", "patterns": []},
}

ALL_FORMATS = {**FORMATS, **ADVISORY_FORMATS}

# Most specific first: "anticipatory bail" before "bail", "cheque bounce notice" before "notice"
FORMAT_ORDER = [
    "cheque_bounce_notice", "section_80_cpc_notice", "anticipatory_bail_application", "bail_application",
    "writ_petition", "consumer_complaint", "rti_application", "written_statement", "plaint",
    "police_complaint", "affidavit", "case_brief", "legal_opinion", "legal_notice",
]
_COMPILED = {k: [re.compile(p, re.IGNORECASE) for p in ALL_FORMATS[k]["patterns"]] for k in FORMAT_ORDER}


def classify_document(query: str) -> str:
    """Sub-classify a drafting request into a document format key (defaults to 'legal_report')."""
    for key in FORMAT_ORDER:
        if any(p.search(query) for p in _COMPILED[key]):
            return key
    return "legal_report"


def describe_format(key: str) -> dict:
    spec = ALL_FORMATS.get(key, ALL_FORMATS["legal_report"])
    return {"key": key, "label": spec["label"], "category": spec["category"], "basis": spec["basis"]}


def field_schema_prompt(key: str) -> str:
    """Field list for the LLM prompt so it fills the right JSON keys for this format."""
    spec = FORMATS.get(key)
    if not spec:
        return ""
    lines = [f'- "{name}": {desc}' for name, desc in spec["fields"].items()]
    return f"Fields for a {spec['label']} (basis: {spec['basis']}). Use these exact keys; lists must be JSON arrays:\n" + "\n".join(lines)


# ----------------------------------------------------------------------------
# Filling fields from free-form Markdown (when the model writes prose, not JSON)
# ----------------------------------------------------------------------------
_FIELD_HINTS = {
    "facts": ["fact", "background", "brief facts", "statement of facts"],
    "grounds": ["ground"],
    "prayer": ["prayer", "relief", "pray"],
    "claims": ["demand", "claim", "relief", "requisition"],
    "legal_provisions": ["provision", "statut", "law", "legal basis"],
    "cause_of_action": ["cause of action"],
    "jurisdiction": ["jurisdiction"],
    "limitation": ["limitation"],
    "valuation": ["valuation", "court fee"],
    "preliminary_objections": ["preliminary", "objection"],
    "parawise_reply": ["para-wise", "parawise", "reply on merits", "merits"],
    "additional_facts": ["additional"],
    "synopsis": ["synopsis", "summary"],
    "list_of_dates": ["list of dates", "chronology", "dates and events"],
    "questions_of_law": ["question"],
    "interim_relief": ["interim"],
    "information_sought": ["information"],
    "statements": ["statement", "affirm", "oath"],
    "deficiency": ["deficiency", "defect", "unfair"],
}


def _split_items(text: str) -> list[str]:
    items = []
    for raw in re.split(r"\n\s*\n|\n(?=\s*(?:[-*•]|\d+[.)]|\([a-z]\)))", text):
        item = re.sub(r"^\s*(?:[-*•]|\d+[.)]|\([a-z]\))\s*", "", raw).strip()
        item = re.sub(r"\*\*(.+?)\*\*", r"\1", item)
        if len(item) > 3:
            items.append(" ".join(item.split()))
    return items


def fill_from_markdown(key: str, text: str, args: dict | None = None) -> dict:
    """Map Markdown sections onto the format's fields by heading keywords; existing args win."""
    spec = FORMATS.get(key)
    args = dict(args or {})
    if not spec:
        return args
    sections, current, buf = [], "", []
    for line in text.splitlines():
        m = re.match(r"^\s*(?:#{1,6}\s+|\*\*)(.+?)(?:\*\*)?\s*:?\s*$", line)
        if m and len(m.group(1)) < 80:
            sections.append((current, "\n".join(buf)))
            current, buf = m.group(1).lower(), []
        else:
            buf.append(line)
    sections.append((current, "\n".join(buf)))

    unused = []
    for heading, body in sections:
        if not body.strip():
            continue
        target = next(
            (f for f in spec["fields"] if not args.get(f) and any(h in heading for h in _FIELD_HINTS.get(f, [f.replace("_", " ")]))),
            None,
        )
        if target:
            args[target] = _split_items(body)
        else:
            unused.extend(_split_items(body))
    # Anything unmatched becomes the facts / statements so nothing written is lost
    for sink in ("facts", "statements", "information_sought", "additional_facts"):
        if sink in spec["fields"] and not args.get(sink) and unused:
            args[sink] = unused
            break
    return args


# ----------------------------------------------------------------------------
# Rendering
# ----------------------------------------------------------------------------
class _Blanks(dict):
    def __missing__(self, key):
        return BLANK


def _as_text(value) -> str:
    if value is None or value == "" or value == []:
        return BLANK
    if isinstance(value, (list, tuple)):
        return "; ".join(_as_text(v) for v in value)
    if isinstance(value, dict):
        return ", ".join(f"{v}" for v in value.values() if v)
    return str(value).strip() or BLANK


def _as_items(value) -> list[str]:
    if value is None or value == "" or value == []:
        return []
    if isinstance(value, (list, tuple)):
        return [_as_text(v) for v in value if _as_text(v) != BLANK]
    return _split_items(str(value)) or [str(value)]


class _Renderer:
    def __init__(self, key: str, args: dict):
        self.key, self.spec = key, FORMATS[key]
        roles = self.spec.get("roles", ("Petitioner", "Respondent"))
        self.values = _Blanks({k: _as_text(v) for k, v in args.items()})
        self.values.setdefault("date", datetime.now().strftime("%d %B %Y"))
        self.values["party_role"] = self.spec.get("verify_as", roles[0])
        verifier_field = "respondent_name" if self.values["party_role"] == roles[1] else "petitioner_name"
        self.values["verifier"] = self.values[verifier_field]
        self.args, self.roles, self.para = args, roles, 0
        self.doc = docx.Document()
        self._setup()

    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Mm(297), Mm(210)  # A4, as required by most Indian courts
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin, sec.right_margin = Inches(1.5), Inches(1)
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = "Times New Roman", Pt(12)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.space_after = Pt(6)

    def fmt(self, template: str) -> str:
        return string.Formatter().vformat(template, (), self.values)

    def para_(self, text="", bold=False, align=None, italic=False, size=None):
        p = self.doc.add_paragraph()
        if align:
            p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
        return p

    def render(self) -> docx.Document:
        for block in self.spec["layout"]:
            self._render_block(block)
        notes = self.spec.get("notes")
        if notes:
            self.doc.add_page_break()
            self.para_("DRAFTER'S CHECKLIST (remove before filing or sending)", bold=True, size=11)
            for n in notes:
                self.para_(f"• {n}", italic=True, size=10)
            self.para_("Generated draft: verify every statutory reference, date and amount against the current bare act and court rules.", italic=True, size=9)
        return self.doc

    def _render_block(self, b: dict):
        t = b["type"]
        if t == "banner":
            self.para_(self.fmt(b["text"]), bold=True, align="center", size=13)
        elif t == "title":
            self.para_(self.fmt(b["text"]).upper(), bold=True, align="center")
        elif t == "heading":
            self.para_(b["text"].upper(), bold=True)
        elif t == "line":
            text = self.fmt(b["text"])
            if b.get("optional") and text.replace(BLANK, "").strip() == "":
                return
            self.para_(text, bold=b.get("bold", False), align=b.get("align", "justify"))
        elif t == "text":
            self.para_(_as_text(self.args.get(b["field"])), align="justify")
        elif t in ("numbered", "bullets", "lettered"):
            items = [self.fmt(i) for i in b["items"]] if "items" in b else _as_items(self.args.get(b["field"]))
            if not items and "field" in b:
                items = [BLANK]
            for i, item in enumerate(items):
                if t == "numbered":
                    self.para += 1
                    label = f"{self.para}. "
                elif t == "lettered":
                    label = f"{string.ascii_uppercase[i % 26]}. "
                else:
                    label = "• "
                self.para_(label + item, align="justify")
        elif t == "cause_title":
            self.para_(self.fmt("IN THE {court_name}").upper(), bold=True, align="center")
            if b.get("extra"):
                self.para_(b["extra"], align="center")
            self.para_(self.fmt("{case_number}"), bold=True, align="center")
            self.para_("IN THE MATTER OF:", bold=True)
            self.para_(self.fmt("{petitioner_name}\n{petitioner_address}") + f"\n…{self.roles[0]}", align="justify")
            self.para_("VERSUS", bold=True, align="center")
            self.para_(self.fmt("{respondent_name}\n{respondent_address}") + f"\n…{self.roles[1]}", align="justify")
        elif t == "prayer":
            self.para_("PRAYER", bold=True)
            self.para_(b["intro"], align="justify")
            items = _as_items(self.args.get("prayer")) or [BLANK]
            for i, item in enumerate(items):
                self.para_(f"({string.ascii_lowercase[i % 26]}) {item};", align="justify")
            self.para_(f"({string.ascii_lowercase[len(items) % 26]}) pass such other and further order(s) as this Hon'ble Court may deem fit and proper in the facts and circumstances of the case.", align="justify")
            self.para_("AND FOR THIS ACT OF KINDNESS, THE " + self.roles[0].split("/")[0].upper() + " SHALL AS IN DUTY BOUND EVER PRAY.", bold=True, align="justify")
        elif t == "verification":
            self.values["last_para"] = str(max(self.para, 1))
            self.para_()
            text = self.fmt(b["text"])
            head, _, rest = text.partition("\n\n")
            self.para_(head, bold=True, align="center")
            self.para_(rest, align="justify")
            self.para_(self.values["party_role"].upper(), bold=True, align="right")
        elif t == "signature":
            self.para_()
            self.para_(self.fmt(b["text"]), align="right")


def missing_fields(key: str, args: dict) -> list[dict]:
    """Fields of a registered format that are still blank in `args` (for a fill-in form).

    Returns [{"name", "description", "is_list"}]; an empty list for advisory formats or complete drafts."""
    spec = FORMATS.get(key)
    if not spec:
        return []
    out = []
    for name, desc in spec["fields"].items():
        value = args.get(name)
        text = _as_text(value)
        if text == BLANK or BLANK in text:
            out.append({"name": name, "description": desc, "is_list": "(list" in desc or isinstance(value, list)})
    return out


def read_docx_preview(path: str) -> tuple[list[tuple[str, bool]], list[str]]:
    """Paragraphs of a generated document as (text, is_bold) plus the drafter's checklist items."""
    d = docx.Document(path)
    body, checklist, in_checklist = [], [], False
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if text.startswith("DRAFTER'S CHECKLIST"):
            in_checklist = True
            continue
        if in_checklist:
            checklist.append(text.lstrip("• ").strip())
        else:
            runs = [r for r in p.runs if r.text.strip()]
            body.append((text, bool(runs) and all(r.bold for r in runs)))
    return body, checklist


def render_format(key: str, args: dict) -> dict:
    """Render a registered format to .docx. Returns {status, file_path, document_type, format}."""
    if key not in FORMATS:
        return {"status": "error", "message": f"Unknown document format '{key}'"}
    doc = _Renderer(key, args).render()
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    slug = re.sub(r"[^A-Za-z0-9]+", "_", FORMATS[key]["label"].split("(")[0]).strip("_")
    path = DOCS_DIR / f"{slug}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    doc.save(str(path))
    return {"status": "success", "file_path": str(path), "document_type": FORMATS[key]["label"], "format": key}
