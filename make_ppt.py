import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette - Professional Deep Legal Navy & Gold Accent
    C_DARK = RGBColor(15, 23, 42)       # Slate 900
    C_CARD = RGBColor(30, 41, 59)       # Slate 800
    C_TEXT = RGBColor(248, 250, 252)    # Slate 50
    C_MUTED = RGBColor(148, 163, 184)   # Slate 400
    C_GOLD = RGBColor(245, 158, 11)     # Amber 500
    C_BLUE = RGBColor(56, 189, 248)     # Sky 400
    C_GREEN = RGBColor(52, 211, 153)    # Emerald 400

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="INDIAN LEGAL AI — DOMAIN-ADAPTED LLM & RAG"):
        # Category / Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = C_GOLD

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = C_TEXT

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    slide_layout = prs.slide_layouts[6] # blank
    s1 = prs.slides.add_slide(slide_layout)
    add_bg(s1)

    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.5))
    tf1 = tbox.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "Indian Legal AI: Domain-Adapted LLM & RAG"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = C_TEXT

    p2 = tf1.add_paragraph()
    p2.text = "Autonomous LoRA Fine-Tuning & 5.65M-Chunk Federated RAG Architecture"
    p2.font.size = Pt(20)
    p2.font.color.rgb = C_GOLD
    p2.space_before = Pt(14)

    p3 = tf1.add_paragraph()
    p3.text = "Indian Legal Intelligence • Qwen3-VL-4B/8B • HPC Distributed Training • LanceDB Vector Retrieval • Tool-Calling Agent"
    p3.font.size = Pt(14)
    p3.font.color.rgb = C_MUTED
    p3.space_before = Pt(30)

    # =========================================================================
    # SLIDE 2: Legal Corpus & Data Engineering
    # =========================================================================
    s2 = prs.slides.add_slide(slide_layout)
    add_bg(s2)
    add_header(s2, "Massive Indian Legal Corpus & Data Pipeline")

    cards = [
        ("📚 5.65 Million Chunks", "Pre-indexed searchable legal chunks across 3 dedicated federated tables.", C_BLUE),
        ("🏛️ High Courts (5.0M)", "Over 5,002,082 passages across Bombay, Delhi, CG, and other High Courts.", C_GOLD),
        ("🇮🇳 Supreme Court (546K)", "546,547 chunks covering historical precedent judgments (1950 - 2025).", C_GREEN),
        ("📜 Central & State Acts (99K)", "99,243 bare act sections including IPC, CrPC, CPC, BNS, BSA, BNSS & NI Act.", C_BLUE),
    ]

    for idx, (title, desc, color) in enumerate(cards):
        x = Inches(0.8 + (idx * 2.95))
        y = Inches(1.8)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.8), Inches(1.8))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        tb = s2.shapes.add_textbox(x + Inches(0.15), y + Inches(0.15), Inches(2.5), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = color

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(11)
        pd.font.color.rgb = C_TEXT
        pd.space_before = Pt(8)

    # Bottom Pipeline Details Box
    bottom_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.9), Inches(11.65), Inches(2.9))
    bottom_box.fill.solid()
    bottom_box.fill.fore_color.rgb = C_CARD
    bottom_box.line.color.rgb = C_MUTED
    
    btb = s2.shapes.add_textbox(Inches(1.0), Inches(4.0), Inches(11.25), Inches(2.6))
    btf = btb.text_frame
    btf.word_wrap = True
    bp1 = btf.paragraphs[0]
    bp1.text = "How the Data was Acquired & Preprocessed for Training & Search:"
    bp1.font.size = Pt(16)
    bp1.font.bold = True
    bp1.font.color.rgb = C_GOLD

    points = [
        "• Raw Legal PDFs & Judgments: Aggregated official court gazettes, IndianKanoon, and legislative registries (~100+ GB raw legal texts).",
        "• Supervised Fine-Tuning (SFT) Dataset: 47,500 curated Q&A pairs (140.5 MB train.jsonl) + 2,500 validation pairs (val.jsonl) filtered below 4,096 tokens.",
        "• Strict Zero-Truncation Strategy: Legal cases exceeding 4,096 tokens were safely dropped rather than cut, ensuring statutory answer integrity was never broken.",
        "• Assistant-Only Masking: Loss was computed strictly on assistant legal reasoning tokens (prompt tokens masked with -100)."
    ]
    for pt in points:
        p = btf.add_paragraph()
        p.text = pt
        p.font.size = Pt(12)
        p.font.color.rgb = C_TEXT
        p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 3: RAG Deep Dive Part 1: Chunking Mechanics & LanceDB Vector Store
    # =========================================================================
    s3 = prs.slides.add_slide(slide_layout)
    add_bg(s3)
    add_header(s3, "RAG Architecture (Part 1): Chunking Strategy & LanceDB Engine")

    rag_cols1 = [
        ("1. Section-Aware Legal Chunking", [
            "• Statutory Boundary Preservation: Bare acts are chunked strictly along Section, Clause, and Sub-clause boundaries so punishments, provisos, and explanations remain unified.",
            "• Sliding Window for Judgments: Complex High Court/Supreme Court cases use 512-token chunks with 64-token overlap to prevent losing continuity across multi-judge benches.",
            "• Provenance Metadata Injection: Every chunk prepends [ACT: IPC] [SECTION: 302] or [COURT: Supreme Court] [CASE: XYZ v. State] directly into text, ensuring embeddings retain contextual authority."
        ], C_BLUE),
        ("2. Dense Embedding Vector Space", [
            "• Model: BAAI/bge-small-en-v1.5 producing compact 384-dimensional dense semantic vectors.",
            "• L2-Normalized Vectors: All chunk embeddings are pre-normalized, converting cosine similarity calculations into blazing-fast inner dot products.",
            "• Completely Offline & Private: Dense vectorization runs locally on CPU/GPU, eliminating API egress costs, rate limits, and client confidentiality breaches."
        ], C_GOLD),
        ("3. LanceDB Columnar Vector Store", [
            "• Serverless Columnar Storage: Disk-backed Apache Arrow format enables zero-copy reads without spinning up external vector server processes (Pinecone/Milvus).",
            "• Sub-10ms ANN Query Speed: Scalable across 5.65M vectors with IVF-PQ (Inverted File Product Quantization) index, operating seamlessly on a consumer laptop.",
            "• Low Memory Footprint: Directly streams vector indices off disk; system RAM consumption remains under 3 GB during active search."
        ], C_GREEN),
    ]

    for idx, (title, bullet_list, col_color) in enumerate(rag_cols1):
        x = Inches(0.8 + (idx * 3.95))
        y = Inches(1.8)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.75), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = col_color
        card.line.width = Pt(1.5)

        tb = s3.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), Inches(3.35), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = col_color

        for b in bullet_list:
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(11)
            p.font.color.rgb = C_TEXT
            p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 4: RAG Deep Dive Part 2: Multi-Stage Federated Retrieval & Re-ranking
    # =========================================================================
    s4 = prs.slides.add_slide(slide_layout)
    add_bg(s4)
    add_header(s4, "RAG Architecture (Part 2): Federated Quotas & Deep Re-ranking")

    rag_cols2 = [
        ("1. Federated Multi-Corpus Quotas", [
            "• The Imbalance Challenge: 5,002,082 High Court passages would easily drown out 99,243 Bare Acts and 546K Supreme Court precedents in a naive global top-K search.",
            "• Dedicated Table Searches: Independent queries dispatched to 'acts', 'supreme_court', and 'high_court' tables with guaranteed candidate quotas.",
            "• Balanced Candidate Pool: Aggregates top candidates across statutory bare laws and precedents, ensuring the legal foundation is never omitted."
        ], C_BLUE),
        ("2. Cross-Encoder Deep Re-ranking", [
            "• Model: ms-marco-MiniLM-L-6-v2 Cross-Encoder evaluated over top-20 merged candidates.",
            "• Full Cross-Attention: Passes (query, passage) jointly through all transformer layers, capturing syntactic and subtle legal relationships that bi-encoders miss.",
            "• Precise Candidate Pruning: Eliminates lexical false positives (e.g. differentiating Anticipatory Bail under Sec 438 CrPC from Regular Bail under Sec 439 CrPC), outputting top-5 relevant chunks."
        ], C_GOLD),
        ("3. Adaptive Router & Tool-Calling Agent", [
            "• Intent Classification: LLM router classifies queries into Statute, Precedent, Grievance, or Current Information.",
            "• Confidence-Driven Web Fallback: If reranker score is below -2.0, Tavily live web search activates automatically to fetch current gazettes or recent amendments.",
            "• Autonomous Agent Integration: Decouples reasoning from rendering to invoke external tools (Legal Notice/Opinion .docx generator, Sec 34 CPC calculator)."
        ], C_GREEN),
    ]

    for idx, (title, bullet_list, col_color) in enumerate(rag_cols2):
        x = Inches(0.8 + (idx * 3.95))
        y = Inches(1.8)
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.75), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = col_color
        card.line.width = Pt(1.5)

        tb = s4.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), Inches(3.35), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = col_color

        for b in bullet_list:
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(11)
            p.font.color.rgb = C_TEXT
            p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 5: LoRA Fine-Tuning on HPC Cluster
    # =========================================================================
    s5 = prs.slides.add_slide(slide_layout)
    add_bg(s5)
    add_header(s5, "HPC Fine-Tuning Execution & Infrastructure")

    left_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0))
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = C_CARD
    left_box.line.color.rgb = C_BLUE

    ltb = s5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.6))
    ltf = ltb.text_frame
    ltf.word_wrap = True
    lp1 = ltf.paragraphs[0]
    lp1.text = "Cluster Hardware & LoRA Parameters"
    lp1.font.size = Pt(18)
    lp1.font.bold = True
    lp1.font.color.rgb = C_BLUE

    cluster_specs = [
        "• Cluster: Plaksha HPC Slurm Cluster (partition: rtx, general QOS).",
        "• GPU Hardware: 1x 48 GB NVIDIA RTX PRO 6000 Blackwell MIG slice (gpu:2g.48gb:1).",
        "• Base Model: Qwen/Qwen3-VL-4B-Instruct (Vision-Language architecture).",
        "• LoRA Hyperparameters: Rank r=32, Alpha=64, Dropout=0.05, rsLoRA enabled.",
        "• Target Modules: All Attention & MLP projections (q, k, v, o, gate, up, down).",
        "• Trainable Parameters: 66,060,288 params (only 1.47% of total 4.5B params!).",
        "• Effective Batch Size: 16 (micro-batch 4, grad accumulation 4)."
    ]
    for cs in cluster_specs:
        p = ltf.add_paragraph()
        p.text = cs
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_TEXT
        p.space_before = Pt(8)

    right_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = C_CARD
    right_box.line.color.rgb = C_GOLD

    rtb = s5.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    rtf = rtb.text_frame
    rtf.word_wrap = True
    rp1 = rtf.paragraphs[0]
    rp1.text = "Self-Chaining & Fault-Tolerant Training"
    rp1.font.size = Pt(18)
    rp1.font.bold = True
    rp1.font.color.rgb = C_GOLD

    training_highlights = [
        "• Automated Self-Chaining: Handles short Slurm wall-times automatically. Stops cleanly 15 mins before timeout, writes checkpoint, and re-submits itself.",
        "• Safe Checkpointing: Rolling checkpoints saved every 200 steps to /scratch, and synced adapter mirrored to persistent /home.",
        "• Full 2-Epoch Convergence: Ran across 47,500 samples in exactly 7h 02m.",
        "• Loss Monotonic Drop: Training loss dropped from 1.949 to 0.592; Eval loss dropped from 0.847 to 0.725 with zero gradient spikes.",
        "• Ready for Edge Inference: Fused LoRA adapter weights take only 264 MB!"
    ]
    for th in training_highlights:
        p = rtf.add_paragraph()
        p.text = th
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_TEXT
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 6: Benchmark Evaluation: Base vs Fine-Tuned
    # =========================================================================
    s6 = prs.slides.add_slide(slide_layout)
    add_bg(s6)
    add_header(s6, "Benchmark Evaluation: Qwen Base vs Fine-Tuned Legal Model")

    # Table on Slide 6
    rows, cols = 8, 5
    table_shape = s6.shapes.add_table(rows, cols, Inches(0.8), Inches(1.8), Inches(11.7), Inches(3.6))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(2.8)
    tbl.columns[1].width = Inches(1.9)
    tbl.columns[2].width = Inches(2.2)
    tbl.columns[3].width = Inches(1.8)
    tbl.columns[4].width = Inches(3.0)

    headers = ["Metric Evaluated", "Base Model (4B)", "Fine-Tuned Legal LoRA", "Winner", "Improvement / Gain"]
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(40, 53, 76)
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_GOLD

    table_data = [
        ("Perplexity (PPL) ↓", "3.47", "2.04", "Fine-Tuned", "41.2% lower uncertainty 🚀"),
        ("Cross-Entropy Loss ↓", "1.2453", "0.7132", "Fine-Tuned", "42.7% lower loss (High confidence)"),
        ("ROUGE-1 (Unigrams) ↑", "27.52%", "31.15%", "Fine-Tuned", "+13.2% relative concept gain"),
        ("ROUGE-2 (Legal Bigrams) ↑", "10.78%", "17.12%", "Fine-Tuned", "+58.8% massive legal phrasing jump 🔥"),
        ("ROUGE-L (Sentence flow) ↑", "18.18%", "22.33%", "Fine-Tuned", "+22.8% structural alignment"),
        ("BLEU-4 (Exact 4-gram) ↑", "2.68", "5.34", "Fine-Tuned", "+99.3% relative gain (2x better)"),
        ("Legal Citation Recall ↑", "66.92%", "66.51%", "Parity", "High citation accuracy on both (~67%)"),
    ]

    for row_idx, row_values in enumerate(table_data, start=1):
        for col_idx, val in enumerate(row_values):
            cell = tbl.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_CARD if row_idx % 2 == 0 else RGBColor(22, 30, 46)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(11)
            p.font.color.rgb = C_GREEN if col_idx == 3 and "Fine-Tuned" in val else C_TEXT
            if col_idx == 4 and ("41" in val or "58" in val or "2x" in val):
                p.font.bold = True
                p.font.color.rgb = C_GOLD

    # Takeaway Callout Box
    c_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.4))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = C_CARD
    c_box.line.color.rgb = C_GREEN

    ctb = s6.shapes.add_textbox(Inches(1.0), Inches(5.65), Inches(11.3), Inches(1.3))
    ctf = ctb.text_frame
    ctf.word_wrap = True
    cp = ctf.paragraphs[0]
    cp.text = "Key Scientific Finding:"
    cp.font.size = Pt(14)
    cp.font.bold = True
    cp.font.color.rgb = C_GREEN

    cp2 = ctf.add_paragraph()
    cp2.text = "The base model answers legal prompts with generic conversational prose. Fine-tuning on 47.5k Indian legal pairs taught the model judicial phrasing: doubling BLEU-4 (+99.3%) and surging ROUGE-2 by +58.8%, while cutting perplexity by 41.2%."
    cp2.font.size = Pt(11.5)
    cp2.font.color.rgb = C_TEXT
    cp2.space_before = Pt(4)

    # Save with fallback in case PowerPoint has an active file lock on the original file
    out_file = "Indian_Legal_AI_Presentation.pptx"
    v2_file = "Indian_Legal_AI_Presentation_v2.pptx"
    try:
        prs.save(out_file)
        print(f"Presentation saved successfully as '{out_file}'!")
    except PermissionError:
        prs.save(v2_file)
        print(f"Original file was open in PowerPoint. Saved updated presentation to '{v2_file}'!")

if __name__ == "__main__":
    create_presentation()
